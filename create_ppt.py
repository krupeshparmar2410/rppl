from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_presentation():
    # Create presentation
    prs = Presentation()
    
    # Define common colors based on the UI
    dark_bg = RGBColor(11, 16, 32)
    cyan = RGBColor(6, 182, 212)
    blue = RGBColor(59, 130, 246)
    white = RGBColor(255, 255, 255)
    gray = RGBColor(148, 163, 184)
    emerald = RGBColor(16, 185, 129)
    red = RGBColor(239, 68, 68)

    # Helper function to apply dark background
    def apply_dark_bg(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = dark_bg

    # Helper function to format titles
    def format_title(title_shape, text, size=40, color=cyan):
        title_shape.text = text
        for paragraph in title_shape.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER
            for run in paragraph.runs:
                run.font.size = Pt(size)
                run.font.color.rgb = color
                run.font.bold = True
                run.font.name = "Arial"

    # Slide 1: Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    apply_dark_bg(slide)
    
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    format_title(title, "TRANSFORMER AI\nHEALTH MONITORING SYSTEM", size=44, color=cyan)
    
    subtitle.text = "AI-powered predictive maintenance and industrial monitoring"
    for p in subtitle.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(24)
            r.font.color.rgb = white

    # Slide 2: The Challenge
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    apply_dark_bg(slide)
    
    format_title(slide.shapes.title, "The Challenge in Power Grids", size=36)
    
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    tf.text = "Power transformers are critical infrastructure."
    
    p = tf.add_paragraph()
    p.text = "Unplanned Failures: Costly downtimes and severe damage."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Manual Inspections: Slow, dangerous, and often too late."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Complex Data: Human operators cannot analyze thousands of sensor parameters in real time."
    p.level = 1
    
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = white
            run.font.size = Pt(24)

    # Slide 3: Our AI Solution
    slide = prs.slides.add_slide(bullet_slide_layout)
    apply_dark_bg(slide)
    
    format_title(slide.shapes.title, "Our Smart Intelligence Platform", size=36)
    
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    tf.text = "An end-to-end Machine Learning pipeline:"
    
    p = tf.add_paragraph()
    p.text = "Real-Time Telemetry: Processing Voltage, Current, OTI, and Power Factors instantly."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Predictive Inference: Classifying health into Normal, Warning, or Critical."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Anomaly Detection: Catching unusual patterns before major faults occur."
    p.level = 1
    
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = white
            run.font.size = Pt(24)

    # Slide 4: Model Performance
    slide = prs.slides.add_slide(bullet_slide_layout)
    apply_dark_bg(slide)
    
    format_title(slide.shapes.title, "Machine Learning Excellence", size=36)
    
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    tf.text = "State-of-the-art predictive modeling:"
    
    p = tf.add_paragraph()
    p.text = "Algorithm: Random Forest & XGBoost Ensembles"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Dataset: 19,000+ Industrial Sensor Readings"
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Feature Set: 32 engineered inputs (e.g., Thermal Stress, Load Level, THD)"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Accuracy: 97.3% Prediction Confidence"
    p.level = 1
    
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = white
            run.font.size = Pt(24)

    # Slide 5: The Dashboard UI
    slide = prs.slides.add_slide(bullet_slide_layout)
    apply_dark_bg(slide)
    
    format_title(slide.shapes.title, "Live Monitoring Dashboard", size=36)
    
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    tf.text = "Sleek, intuitive, and modern Web Interface:"
    
    p = tf.add_paragraph()
    p.text = "Dark Mode aesthetics with glassmorphism UI."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Live CSV Uploads: Instantly evaluates batch sensor data."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Visual Data: Real-time probability bars and metrics."
    p.level = 1
    
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = white
            run.font.size = Pt(24)

    # Slide 6: Conclusion
    slide = prs.slides.add_slide(title_slide_layout)
    apply_dark_bg(slide)
    
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    format_title(title, "Thank You", size=50, color=cyan)
    
    subtitle.text = "Ready to deploy for next-generation power grids."
    for p in subtitle.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(28)
            r.font.color.rgb = white

    # Save presentation
    prs.save("Transformer_AI_Presentation.pptx")
    print("Presentation created successfully as 'Transformer_AI_Presentation.pptx'")

if __name__ == "__main__":
    create_presentation()
