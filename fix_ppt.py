from pptx import Presentation

def fix_ppt(file_path, out_path):
    prs = Presentation(file_path)
    
    # 1. Slide 12: Move content up by 800,000 units
    slide_12 = prs.slides[11]
    for i in range(7, len(slide_12.shapes)):
        shape = slide_12.shapes[i]
        if hasattr(shape, 'top'):
            shape.top = shape.top - 800000

    # 2. Slide 13: Adjust column spacing
    slide_13 = prs.slides[12]
    
    # Col 1
    slide_13.shapes[9].left = 772000
    slide_13.shapes[9].width = 1600000
    slide_13.shapes[10].left = 822000
    slide_13.shapes[10].width = 1500000
    slide_13.shapes[11].left = 822000
    slide_13.shapes[11].width = 1500000
    
    # Col 2
    slide_13.shapes[12].left = 2772000
    slide_13.shapes[12].width = 1600000
    slide_13.shapes[13].left = 2822000
    slide_13.shapes[13].width = 1500000
    slide_13.shapes[14].left = 2822000
    slide_13.shapes[14].width = 1500000
    
    # Col 3
    slide_13.shapes[15].left = 4772000
    slide_13.shapes[15].width = 1600000
    slide_13.shapes[16].left = 4822000
    slide_13.shapes[16].width = 1500000
    slide_13.shapes[17].left = 4822000
    slide_13.shapes[17].width = 1500000
    
    # Col 4
    slide_13.shapes[18].left = 6772000
    slide_13.shapes[18].width = 1600000
    slide_13.shapes[19].left = 6822000
    slide_13.shapes[19].width = 1500000
    slide_13.shapes[20].left = 6822000
    slide_13.shapes[20].width = 1500000

    # 3. Slide 14: Fix "Thank" and "You"
    slide_14 = prs.slides[13]
    shape_thank = slide_14.shapes[3]
    shape_you = slide_14.shapes[4]
    
    # We append " You" to the first paragraph of shape 3
    if shape_thank.has_text_frame:
        p = shape_thank.text_frame.paragraphs[0]
        # To preserve formatting, copy font from the first run
        if len(p.runs) > 0:
            run = p.add_run()
            run.text = " You"
            run.font.name = p.runs[0].font.name
            run.font.size = p.runs[0].font.size
            run.font.bold = p.runs[0].font.bold
            run.font.italic = p.runs[0].font.italic
            if p.runs[0].font.color.type is not None:
                try:
                    run.font.color.rgb = p.runs[0].font.color.rgb
                except:
                    pass
        else:
            shape_thank.text = "Thank You"
            
    # Clear shape 4
    if shape_you.has_text_frame:
        shape_you.text = ""

    prs.save(out_path)
    print(f"Saved fixed presentation to {out_path}")

if __name__ == "__main__":
    fix_ppt("RPPL_AI_Presentation-1.pptx", "RPPL_AI_Presentation-1.pptx")
