from pptx import Presentation
import sys

def analyze_ppt(file_path):
    prs = Presentation(file_path)
    with open("output.txt", "w", encoding="utf-8") as f:
        f.write(f"Total slides: {len(prs.slides)}\n")
        
        for i in [11, 12, 13]: # Slide 12, 13, 14 (0-indexed)
            if i >= len(prs.slides):
                f.write(f"Slide {i+1} does not exist.\n")
                continue
                
            slide = prs.slides[i]
            f.write(f"\n--- Slide {i+1} ---\n")
            for j, shape in enumerate(slide.shapes):
                f.write(f"Shape {j}: name='{shape.name}', type={shape.shape_type}\n")
                if hasattr(shape, 'text'):
                    f.write(f"Text preview: {repr(shape.text[:100])}\n")
                if hasattr(shape, 'left'):
                    f.write(f"Left: {shape.left}, Top: {shape.top}, Width: {shape.width}, Height: {shape.height}\n")

if __name__ == "__main__":
    analyze_ppt("RPPL_AI_Presentation-1.pptx")
